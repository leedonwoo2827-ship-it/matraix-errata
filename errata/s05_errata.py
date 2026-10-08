"""s05 정오표 — 내부 검수용 엑셀 + 독자 게시용 HTML. 모델을 부르지 않는다.

★ 사람의 판정은 엑셀 「정오표」 시트의 「판정」·「확정 메모」 열에 적는다(반영 · 보류 · 거절).
  다시 돌릴 때 **먼저 그 열을 읽어 decisions.json 에 옮겨 두고** 새 엑셀에 되돌려 적는다 — 다시 돌려도 판정이 사라지지 않는다.
  열쇠는 fid(uid·잘못·바름·유형의 해시)라서, 같은 지적이면 번호가 바뀌어도 판정이 따라간다.
★ 게시 HTML 에는 판정이 「반영」인 행만 나간다. 사람이 확정하지 않은 것은 독자에게 나가지 않는다.
★ 엑셀이 열려 있어 덮을 수 없으면 시각을 붙인 새 이름으로 저장한다(판정은 decisions.json 에 이미 옮겨져 있다).
"""
from __future__ import annotations

import collections
import datetime as dt
import html
import time
from pathlib import Path
from typing import Dict, List

from errata.config import Job
from errata.ingest import PRODUCTS
from errata.util import log, norm, read_json, read_jsonl, warn, write_json

COLS = [("번호", 6), ("판정", 9), ("제안", 11), ("산물", 13), ("위치", 40), ("잘못 (AS IS)", 50), ("바름 (TO BE)", 50),
        ("유형", 9), ("심각도", 7), ("확신", 7), ("재검증", 22), ("출처", 8), ("검수자", 10), ("근거", 60), ("파급 위치", 50),
        ("웹확인 검색어", 30), ("확정 메모", 30), ("파일", 40), ("uid", 24), ("규칙", 12), ("fid", 12)]
DECISIONS = ["반영", "보류", "거절"]


def _xlsx_path(job: Job) -> Path:
    return job.p("05_정오표", "정오표_내부.xlsx")


def read_decisions(job: Job) -> Dict[str, dict]:
    """엑셀(사람이 고친 판) → decisions.json. 엑셀이 없거나 못 읽으면 저장된 것만."""
    dpath = job.p("05_정오표", "decisions.json")
    dec: Dict[str, dict] = read_json(dpath, {}) or {}
    xp = _xlsx_path(job)
    if xp.exists():
        try:
            import openpyxl
            wb = openpyxl.load_workbook(xp, read_only=True, data_only=True)
            ws = wb["정오표"]
            rows = ws.iter_rows(values_only=True)
            head = [str(h or "") for h in next(rows)]
            ix = {h: i for i, h in enumerate(head)}
            for r in rows:
                fid = r[ix["fid"]] if "fid" in ix else None
                if not fid:
                    continue
                d = (r[ix["판정"]] or "").strip() if r[ix["판정"]] else ""
                memo = (r[ix["확정 메모"]] or "") if "확정 메모" in ix else ""
                to_be = r[ix["바름 (TO BE)"]] if "바름 (TO BE)" in ix else None
                if d or memo:
                    dec[str(fid)] = {"판정": d, "메모": str(memo), "바름": str(to_be or ""), "at": time.time()}
            wb.close()
        except Exception as e:  # noqa: BLE001
            warn(f"엑셀 판정을 읽지 못했습니다({e}) — 저장된 decisions.json 만 씁니다")
    write_json(dpath, dec)
    return dec


def _pdf_pages(path: str, cache: Dict[str, List[str]]) -> List[str]:
    if path not in cache:
        try:
            import pymupdf
            with pymupdf.open(path) as d:
                cache[path] = [norm(pg.get_text()) for pg in d]
        except Exception:  # noqa: BLE001
            cache[path] = []
    return cache[path]


def _page_of(f: dict, units: Dict[str, dict], cache) -> str:
    u = units.get(f["uid"]) or {}
    pdf = (u.get("x") or {}).get("pdf") or ""
    if not pdf:
        return ""
    pages = _pdf_pages(pdf, cache)
    for probe in (norm(f["as_is"])[:24], norm(u.get("title", ""))[:16]):
        if len(probe) >= 8:
            for i, t in enumerate(pages, 1):
                if probe in t:
                    return f"p.{i}"
    return ""


def _write_xlsx(job: Job, rows: List[dict], panel: dict, prog: dict, counts: dict) -> Path:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = openpyxl.Workbook()
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    head_font = Font(color="FFFFFF", bold=True)
    sev_fill = {"상": PatternFill("solid", fgColor="F8D7DA"), "중": PatternFill("solid", fgColor="FFF3CD")}
    wrap = Alignment(wrap_text=True, vertical="top")

    def sheet(ws, data: List[dict]):
        ws.append([c for c, _ in COLS])
        for i, (_, w) in enumerate(COLS, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
            ws.cell(1, i).fill, ws.cell(1, i).font = head_fill, head_font
        for r in data:
            ws.append([r["번호"], r["판정"], r["suggest"], PRODUCTS.get(r["product"], r["product"]), r["loc_full"],
                       r["as_is"], r["to_be"], r["type"], r["severity"], r["confidence"],
                       (lambda v: f"{v.get('verdict', '')} — {v.get('reason', '')}" if v else "")(r.get("verify")), r["origin"], r.get("reviewer", ""),
                       r["reason"], "\n".join(r.get("spread") or []), r.get("web_query", ""), r["메모"], r["file"],
                       r["uid"], r.get("rule", ""), r["fid"]])
            row = ws.max_row
            for c in range(1, len(COLS) + 1):
                ws.cell(row, c).alignment = wrap
            if r["severity"] in sev_fill:
                ws.cell(row, 9).fill = sev_fill[r["severity"]]
            if (r.get("verify") or {}).get("verdict") == "반대":
                ws.cell(row, 11).fill = PatternFill("solid", fgColor="E2E3E5")
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions
        dv = DataValidation(type="list", formula1='"' + ",".join(DECISIONS) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        if ws.max_row > 1:
            dv.add(f"B2:B{ws.max_row}")

    # 요약
    ws = wb.active
    ws.title = "요약"
    now = dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    ws.append([f"{job.cfg.get('label', job.name)} — 내부 정오표", now])
    ws.append(["판정 방법", "「정오표」 시트의 판정 열에 반영·보류·거절을 고르고 저장 → run.bat errata (또는 웹앱 「정오표」) → 게시 HTML 에 반영만 나간다"])
    ws.append(["검토 진행", f"배치 {prog.get('done', 0)}/{prog.get('total', 0)} · 실패 {len(prog.get('failed') or [])}"])
    ws.append([])
    ws.append(["산물", "단위 수", "지적", "상", "중", "하", "반영", "보류", "거절", "미판정"])
    for p in "QVSDLBT":
        rs = [r for r in rows if r["product"] == p]
        ws.append([PRODUCTS[p], counts.get(p, 0), len(rs), *(sum(1 for r in rs if r["severity"] == s) for s in "상중하"),
                   *(sum(1 for r in rs if r["판정"] == d) for d in DECISIONS), sum(1 for r in rs if not r["판정"])])
    ws.append([])
    ws.append(["규칙(기계) / 검토단", "건수", "예"])
    byrule = collections.defaultdict(list)
    for r in rows:
        byrule[r.get("rule") or r["origin"]].append(r)
    for k, rs in sorted(byrule.items(), key=lambda kv: -len(kv[1])):
        ws.append([k, len(rs), f"{rs[0]['loc_full']} — {rs[0]['as_is'][:60]} → {rs[0]['to_be'][:60]}"])
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width, ws.column_dimensions["C"].width = 28, 14, 110
    for c in ws[1]:
        c.font = Font(bold=True, size=13)

    sheet(wb.create_sheet("정오표"), rows)
    for p in "QVSDLBT":
        rs = [r for r in rows if r["product"] == p]
        if rs:
            sheet(wb.create_sheet(f"{p}_{PRODUCTS[p]}"[:31]), rs)
    web = [r for r in rows if r.get("needs_web")]
    ws = wb.create_sheet("웹확인 필요")
    ws.append(["번호", "위치", "잘못", "바름(제안)", "검색어", "검색 링크", "근거", "fid"])
    for r in web:
        q = r.get("web_query") or r["as_is"][:40]
        ws.append([r["번호"], r["loc_full"], r["as_is"], r["to_be"], q,
                   "https://www.google.com/search?q=" + html.escape(q).replace(" ", "+"), r["reason"], r["fid"]])
    for col, w in zip("ABCDEFGH", (6, 40, 50, 50, 30, 40, 60, 12)):
        ws.column_dimensions[col].width = w
    ws = wb.create_sheet("검수단 명부")
    ws.append(["번호", "구분", "자리", "가명", "묻는 것/층", "약력"])
    for p in panel.get("panel") or []:
        ws.append([p.get("번호"), "검수진" if p["role"] == "author" else "수험생", p.get("view") or p.get("seat"),
                   p.get("가명"), p.get("asks") or p.get("desc"), p.get("약력")])
    ws.append([])
    ws.append(["재현", f"MatrAIx Persona-1M · 씨앗 {panel.get('seed')} · {', '.join(panel.get('shards') or [])}"])
    ws.append(["한계", "가상의 응답자다. 실제 전문가·수험생 검수를 대체하지 않는다. 확정은 사람이 한다."])
    for col, w in zip("ABCDEF", (8, 8, 16, 10, 50, 100)):
        ws.column_dimensions[col].width = w

    xp = _xlsx_path(job)
    xp.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(xp)
        return xp
    except PermissionError:
        alt = xp.with_name(f"정오표_내부_{dt.datetime.now():%m%d-%H%M}.xlsx")
        wb.save(alt)
        warn(f"엑셀이 열려 있어 덮지 못했습니다 → {alt.name} 로 저장 (판정은 decisions.json 에 보존됨)")
        return alt


CSS = """
:root{--bg:#fff;--fg:#1b1f24;--mut:#5b6573;--line:#d9dee5;--head:#f3f5f8;--acc:#1f5fbf;--bad:#b42318;--good:#067647}
@media (prefers-color-scheme:dark){:root:not([data-theme="light"]){--bg:#111418;--fg:#e6e9ee;--mut:#9aa4b2;--line:#2b313a;--head:#1a1f26;--acc:#7fb0ff;--bad:#ff8a80;--good:#6fdc9a}}
:root[data-theme="dark"]{--bg:#111418;--fg:#e6e9ee;--mut:#9aa4b2;--line:#2b313a;--head:#1a1f26;--acc:#7fb0ff;--bad:#ff8a80;--good:#6fdc9a}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--fg);font:15px/1.6 "Pretendard","Noto Sans KR",system-ui,sans-serif}
main{max-width:1080px;margin:0 auto;padding:32px 16px 64px}h1{font-size:26px;margin:0 0 4px}.sub{color:var(--mut);margin:0 0 28px}
h2{font-size:18px;margin:36px 0 10px;padding-bottom:6px;border-bottom:2px solid var(--acc)}
.tbl{overflow-x:auto}table{width:100%;border-collapse:collapse;font-size:14px}th,td{border-bottom:1px solid var(--line);padding:8px 10px;text-align:left;vertical-align:top}
th{background:var(--head);font-weight:600;white-space:nowrap}td.loc{white-space:nowrap;color:var(--mut)}td.bad{color:var(--bad)}td.good{color:var(--good)}
.empty{padding:24px;border:1px dashed var(--line);border-radius:8px;color:var(--mut)}footer{margin-top:40px;color:var(--mut);font-size:13px}
@media (max-width:640px){td.loc{white-space:normal}}
"""


def _write_html(job: Job, rows: List[dict]) -> Path:
    pub = [r for r in rows if r["판정"] == "반영"]
    label = job.pack().get("label", job.name)
    now = dt.datetime.now().strftime("%Y년 %m월 %d일")
    parts = [f"<!doctype html><html lang='ko'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>"
             f"<title>{html.escape(label)} 정오표</title><style>{CSS}</style></head><body><main>",
             f"<h1>{html.escape(label)} 정오표</h1><p class='sub'>최종 갱신 {now} · 확정 {len(pub)}건 · "
             f"교재·강의·영상에서 바로잡은 내용입니다.</p>"]
    if not pub:
        parts.append("<div class='empty'>확정 대기 중입니다. 검수를 마친 항목이 이곳에 게시됩니다.</div>")
    for p in "BLDVQST":
        rs = [r for r in pub if r["product"] == p]
        if not rs:
            continue
        parts.append(f"<h2>{html.escape(PRODUCTS[p])}</h2><div class='tbl'><table><thead><tr><th>위치</th><th>잘못</th><th>바름</th></tr></thead><tbody>")
        for r in rs:
            to_be = r.get("바름_확정") or r["to_be"]
            parts.append(f"<tr><td class='loc'>{html.escape(r['loc_full'])}</td><td class='bad'>{html.escape(r['as_is'])}</td>"
                         f"<td class='good'>{html.escape(to_be)}</td></tr>")
        parts.append("</tbody></table></div>")
    parts.append("<footer>오류를 발견하시면 알려 주십시오. 확인 후 이 정오표에 반영합니다.</footer></main></body></html>")
    out = job.p("05_정오표", "정오표_게시.html")
    out.write_text("".join(parts), encoding="utf-8")
    return out


def run(job: Job, quiet: bool = False) -> dict:
    merged = read_jsonl(job.p("04_수용", "merged.jsonl"))
    if not merged and not quiet:
        warn("04_수용/merged.jsonl 이 비었습니다 — merge 를 먼저 돌리십시오")
    dec = read_decisions(job)
    units = {u["uid"]: u for u in read_jsonl(job.p("00_대상", "units.jsonl"))}
    counts = collections.Counter(u["product"] for u in units.values())
    cache: Dict[str, List[str]] = {}
    rows = []
    for i, f in enumerate(merged, 1):
        r = dict(f)
        r["번호"] = i
        d = dec.get(f["fid"]) or {}
        r["판정"] = d.get("판정", "") if d.get("판정") in DECISIONS else ""
        r["메모"] = d.get("메모", "")
        if d.get("바름") and d["바름"] != f["to_be"]:
            r["바름_확정"] = d["바름"]            # 사람이 엑셀에서 바름을 고쳤으면 그것을 쓴다
            r["to_be"] = d["바름"]
        pg = _page_of(f, units, cache) if f["product"] == "B" else ""
        r["loc_full"] = (f["loc"].replace("교재 ", f"교재 {pg} · ", 1) if pg else f["loc"])
        rows.append(r)
    panel = read_json(job.p("01_선정", "패널.json"), {}) or {}
    prog = read_json(job.p("progress.json"), {}) or {}
    xp = _write_xlsx(job, rows, panel, prog, counts)
    hp = _write_html(job, rows)
    n_pub = sum(1 for r in rows if r["판정"] == "반영")
    if not quiet:
        log(f"  → {xp}")
        log(f"  → {hp}  (게시 {n_pub}건 · 판정 {sum(1 for r in rows if r['판정'])}/{len(rows)})")
    return {"rows": len(rows), "published": n_pub, "xlsx": str(xp), "html": str(hp)}
