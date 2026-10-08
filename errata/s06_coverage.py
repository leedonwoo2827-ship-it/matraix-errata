"""s06 검수 기록 — 「무엇을 봤나」를 단위마다 남기는 엑셀. 모델을 부르지 않는다. `check` 단계가 같이 뽑는다.

나중에 산물에서 오류가 나왔을 때 역추적용이다: 그 단위가 어느 배치에서·언제·어떤 기준으로 검토됐고,
결과가 이상없음이었는지·지적이었는지·검토를 안 거쳤는지(미검토·언급 없음·기계검사만)를 한 줄로 찾는다.

    06_검수표/검수기록.xlsx
      읽는 법     열·상태 뜻
      검수 기준   산물 × 단계(기계검사·검수단 검토·재검증) × 항목 — 무엇을 보는가
      배치 기록   검토 배치마다 단위 수·시각·지적·이상없음·언급 없음·프롬프트 기록
      단위별 기록 단위 한 줄 = 어떤 검사를 거쳤고 결과가 무엇인가

★ 검수 기준은 코드에서 읽는다(s02 docstring 규칙 표 · s03 FOCUS · 프롬프트 「진행」 · persona.yaml 자리) — 따로 적어 두면 어긋난다.
"""
from __future__ import annotations

import collections
import datetime as dt
import re
import time
from pathlib import Path
from typing import Dict, List

from errata.config import Job
from errata.ingest import PRODUCTS
from errata.util import log, read_json, read_jsonl, warn

ST_CLEAN, ST_HIT, ST_SILENT, ST_TODO, ST_FAIL, ST_NONE = "이상없음", "지적 있음", "언급 없음", "미검토", "실패", "검토 대상 아님"
ST_FILL = {ST_HIT: "FFF3CD", ST_SILENT: "F8D7DA", ST_TODO: "F8D7DA", ST_FAIL: "F8D7DA", ST_NONE: "E2E3E5"}


def machine_rules() -> List[tuple]:
    """s02 docstring 의 규칙 표 → [(규칙, 산물 접두 목록, 설명)]."""
    from errata import s02_machine
    out = []
    for line in (s02_machine.__doc__ or "").splitlines():
        m = re.match(r"^\s{4}(\S+)\s{2,}(.+)$", line)
        if not m:
            continue
        names = m.group(1).split("/")
        prods = sorted({n.split("-")[0] for n in names})
        out.append((m.group(1), prods, m.group(2).strip()))
    return out


def _rule_applies(prods: List[str], p: str) -> bool:
    return "*" in prods or p in prods


def _q_steps() -> str:
    t = (Path(__file__).resolve().parent / "prompts" / "03_문항.md").read_text(encoding="utf-8")
    m = re.search(r"## 진행[^\n]*\n(.*?)\n## ", t, re.S)
    return re.sub(r"[`*]", "", m.group(1)).strip() if m else ""


def _criteria(job: Job, units_by_p: Dict[str, int], batched: set, counts: collections.Counter) -> List[list]:
    import yaml
    from errata.s03_review import FOCUS
    from errata.s03b_verify import HEAVY_TYPES
    pcfg = yaml.safe_load((job.pack_dir / "persona.yaml").read_text(encoding="utf-8")) or {}
    rows = []
    rules = machine_rules()
    for p in "QVSDLBT":
        if not units_by_p.get(p):
            continue
        for name, prods, desc in rules:
            if _rule_applies(prods, p):
                n = sum(v for (pp, r), v in counts.items() if pp == p and r in name.replace("*", p).split("/"))
                rows.append([PRODUCTS[p], "① 기계검사", name.replace("*", p), desc, "모든 단위", n])
        if p in batched:
            what = _q_steps() if p == "Q" else FOCUS.get(p, "")
            if p == "Q":
                what += "\n(문제풀이 영상 V 의 해설 낭독도 이 배치에서 함께 본다)"
            rows.append([PRODUCTS[p], "② 검수단 검토", "review", what, f"배치마다 1회 호출 · 단위 {units_by_p[p]}",
                         counts.get((p, "review"), 0) + counts.get((p, "review-solve"), 0)])
            rows.append([PRODUCTS[p], "③ 재검증", "verify",
                         f"상 전부 + 중 가운데 {'·'.join(sorted(HEAVY_TYPES))}을 다른 호출에서 독립 재판정(동의·수정동의·반대·불확실)",
                         "무거운 지적만", ""])
        else:
            rows.append([PRODUCTS[p], "② 검수단 검토", "—",
                         "검수단 검토를 거치지 않는다 — 기계검사만. "
                         + ("낭독 내용은 원천 문항(Q) 검토에서 해설 낭독으로 함께 본다." if p == "V"
                            else "내용이 롱폼(L)·문항(Q)과 같아 그쪽 검토로 갈음한다." if p == "D" else ""), "", ""])
    for s in (pcfg.get("author") or {}).get("seats") or []:
        rows.append(["(검토 산물 공통)", "② 검수단 관점", f"{s.get('id')} {s.get('view')}", s.get("asks", ""), "검수진", ""])
    for s in (pcfg.get("reviewer") or {}).get("strata") or []:
        rows.append(["(검토 산물 공통)", "② 검수단 관점", f"수험생 {s.get('id')}", s.get("desc", ""), f"{s.get('n')}명", ""])
    for r in pcfg.get("require") or []:
        rows.append(["(검토 산물 공통)", "② 검토 규칙", "require", r, "", ""])
    return rows


def _table(ws, cols, data, wrap_cols=(), status_col=None, freeze="A2") -> None:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    head_fill, head_font = PatternFill("solid", fgColor="1F3A5F"), Font(color="FFFFFF", bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    ws.append([c for c, _ in cols])
    for i, (_, w) in enumerate(cols, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        ws.cell(1, i).fill, ws.cell(1, i).font = head_fill, head_font
    for r in data:
        ws.append(r)
        row = ws.max_row
        for c in wrap_cols:
            ws.cell(row, c).alignment = wrap
        if status_col and r[status_col - 1] in ST_FILL:
            ws.cell(row, status_col).fill = PatternFill("solid", fgColor=ST_FILL[r[status_col - 1]])
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions


def write_roster(job: Job, by_rev: collections.Counter, batched: set, units_by_p: collections.Counter) -> Path:
    """01_선정/패널명부.xlsx — 명부 + 산물별 역할(pack/<팩>/roles.yaml). 「지적」은 그 사람 이름으로 나온 검토 지적 수일 뿐,
    옳고 그름·반영 여부와는 별개다."""
    import openpyxl
    import yaml
    rp = job.pack_dir / "roles.yaml"
    rcfg = (yaml.safe_load(rp.read_text(encoding="utf-8")) or {}) if rp.exists() else {}
    roles = rcfg.get("roles") or {}
    rprods = [p for p in (rcfg.get("products") or list("QVSDLBT")) if units_by_p.get(p)]
    pj = read_json(job.p("01_선정", "패널.json"), {}) or {}
    panel = pj.get("panel") or []
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "명부"
    _table(ws, [("번호", 6), ("구분", 8), ("자리·층", 16), ("가명", 9), ("묻는 것 / 층", 50), ("약력", 110)],
           [[m.get("번호"), "검수진" if m.get("role") == "author" else "수험생", f"{m.get('seat')} {m.get('view') or ''}".strip(),
             m.get("가명"), m.get("asks") or m.get("desc"), m.get("약력")] for m in panel], wrap_cols=(5, 6))
    ws.append([])
    ws.append(["재현", "", f"MatrAIx Persona-1M · 씨앗 {pj.get('seed')} · {', '.join(pj.get('shards') or [])}"])
    ws.append(["한계", "", "가상의 응답자다. 실제 전문가·수험생 검수를 대체하지 않는다."])

    cols = [("번호", 6), ("가명", 8), ("구분", 7), ("자리·층", 14)]
    for p in rprods:
        cols += [(f"{PRODUCTS[p]} 역할", 34), ("지적", 6)]
    rows = []
    for m in panel:
        r = roles.get(m.get("seat")) or {}
        row = [m.get("번호"), m.get("가명"), "검수진" if m.get("role") == "author" else "수험생",
               f"{m.get('seat')} {m.get('view') or ''}".strip()]
        for p in rprods:
            row += [r.get(p, ""), by_rev.get((p, m.get("번호")), 0) if p in batched else ""]
        rows.append(row)
    ws = wb.create_sheet("산물별 역할")
    _table(ws, cols, rows, wrap_cols=tuple(5 + 2 * i for i in range(len(rprods))), freeze="E2")
    ws.append([])
    ws.append(["※", "", "", "", f"역할은 pack/{job.pack_dir.name}/roles.yaml 에서 읽는다 — 고도화할 때 그 파일만 고친다. "
               "「지적」은 그 사람 이름으로 나온 검토 지적 수일 뿐, 옳고 그름·반영 여부와는 별개다. "
               "문제풀이 영상·덱은 검수단 검토를 거치지 않는다(기계검사 + 원천 검토로 갈음)."])

    out = job.p("01_선정", "패널명부.xlsx")
    try:
        wb.save(out)
    except PermissionError:
        out = out.with_name(f"패널명부_{dt.datetime.now():%m%d-%H%M}.xlsx")
        wb.save(out)
        warn(f"패널명부.xlsx 가 열려 있어 덮지 못했습니다 → {out.name} 로 저장")
    log(f"  패널명부 → {out.name}")
    return out


def _chapter(u: dict, subj_of_slug: Dict[str, int]):
    """단위 → (과목, 장). 산물마다 배치 id 에 과목·장이 들어 있다: L s1-03 · B:01:03 · S:<slug>:03 · T:q1-03."""
    p, b, x = u["product"], u.get("batch") or "", u.get("x") or {}
    try:
        if p == "L":
            m = re.match(r"s(\d+)-(\d+)", x.get("section") or "")
            return (int(m.group(1)), int(m.group(2))) if m else None
        if p == "B":
            _, vol, ch = b.split(":")
            return int(vol), max(int(ch), 1)          # 00 = 권 머리말 → 첫 장에 붙인다
        if p == "S":
            _, slug, ch = b.split(":")
            return subj_of_slug.get(slug), int(ch)
        if p == "T":
            m = re.match(r"T:q(\d+)-(\d+)", b)
            return (int(m.group(1)), int(m.group(2))) if m else None
    except (ValueError, AttributeError):
        return None
    return None


def write_grid(wb, job: Job, units: List[dict], reviews: Dict[str, dict], links: Dict[str, List[str]]) -> None:
    """「진행표」 — 과목 × 강(롱폼 L01~) × 산물. 칸 = 본 것/전체, 다 보면 진하게·일부면 옅게 칠한다.
    문제집(Q)은 연결된 요약·롱폼 항목의 장으로, 문제풀이 영상(V)은 그 문항을 따라간다(V 는 Q 검토에서 함께 본다).
    예상 완료 = 남은 배치 수 × 지금까지 배치당 평균 초(사용 한도 대기는 넣지 않았다)."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from errata.s03_review import _batches
    U = {u["uid"]: u for u in units}
    subj_of_slug = {}
    for u in units:
        if u["product"] == "S" and (u.get("x") or {}).get("subject"):
            subj_of_slug[(u.get("batch") or "::").split(":")[1]] = int(u["x"]["subject"])
    ch_of: Dict[str, tuple] = {}
    lec_of_ch: Dict[tuple, str] = {}
    for u in units:
        c = _chapter(u, subj_of_slug)
        if c and c[0]:
            ch_of[u["uid"]] = c
            if u["product"] == "L":
                lec_of_ch.setdefault(c, (u.get("x") or {}).get("lecture"))
    for q, ds in links.items():                       # Q → 연결된 S·L 의 장 중 가장 많은 것
        cs = collections.Counter(ch_of[d] for d in ds if d in ch_of and d[0] in "SL")
        if cs and q in U:
            ch_of[q] = cs.most_common(1)[0][0]
            for d in ds:
                if d.startswith("V:"):
                    ch_of.setdefault(d, ch_of[q])
    q_of_v = {d: q for q, ds in links.items() for d in ds if d.startswith("V:")}

    # 예상 완료
    order = (job.get("review.order") or list("QLBST"))
    pend = [b for b, _ in _batches(units, order, []) if b not in reviews]
    secs = [d.get("sec") or 0 for d in reviews.values() if d.get("sec")]
    avg = (sum(secs) / len(secs)) if secs else 60
    eta = {b: time.time() + avg * (i + 1) for i, b in enumerate(pend)}

    def done(u: dict) -> bool:
        if u["product"] == "V":
            q = U.get(q_of_v.get(u["uid"], ""))
            return bool(q and q.get("batch") in reviews)
        return u.get("batch") in reviews

    cols = [("L", "롱폼 강의"), ("B", "교재"), ("S", "요약노트"), ("Q", "문제집"), ("V", "문제집 영상"), ("T", "쇼츠")]
    cols = [(p, n) for p, n in cols if any(x["product"] == p for x in units)]
    cell: Dict[tuple, List[dict]] = collections.defaultdict(list)
    unplaced = collections.Counter()
    for u in units:
        if u["product"] not in dict(cols):
            continue
        c = ch_of.get(u["uid"])
        lec = lec_of_ch.get(c) if c else None
        if not lec:
            unplaced[u["product"]] += 1
            continue
        cell[(lec, u["product"])].append(u)

    lec_subj: Dict[str, int] = {}
    lec_chs: Dict[str, List[tuple]] = collections.defaultdict(list)
    for c, lec in sorted(lec_of_ch.items()):
        lec_subj.setdefault(lec, c[0])
        lec_chs[lec].append(c)
    subjects = job.pack().get("subjects") or {}

    ws = wb.create_sheet("진행표", 0)
    full, part, none_ = PatternFill("solid", fgColor="8EA9DB"), PatternFill("solid", fgColor="DDEBF7"), PatternFill("solid", fgColor="F2F2F2")
    bold, center = Font(bold=True), Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.append([f"진행표 — {job.cfg.get('label', job.name)}  (갱신 {dt.datetime.now():%m-%d %H:%M})"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["과목", "강", "장"] + [n for _, n in cols] + ["이 강 예상 완료"])
    for i in range(1, len(cols) + 5):
        ws.cell(2, i).font, ws.cell(2, i).alignment = bold, center
    widths = [14, 6, 16] + [12] * len(cols) + [16]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl_col(i)].width = w
    start, prev = 3, None
    for lec in sorted(lec_subj, key=lambda k: (lec_subj[k], k)):
        s = lec_subj[lec]
        chs = lec_chs[lec]
        row = [f"{s}과목 {subjects.get(s, '')}".strip() if s != prev else "", f"{int(lec[1:])}강",
               f"{chs[0][0]}-{chs[0][1]:02d}" + (f" ~ {chs[-1][0]}-{chs[-1][1]:02d}" if len(chs) > 1 else "")]
        etas = []
        for p, _ in cols:
            us = cell.get((lec, p), [])
            n = sum(1 for u in us if done(u))
            row.append(f"{n}/{len(us)}" if us else "—")
            for u in us:
                b = u.get("batch") if p != "V" else (U.get(q_of_v.get(u["uid"], "")) or {}).get("batch")
                if b in eta:
                    etas.append(eta[b])
        row.append(dt.datetime.fromtimestamp(max(etas)).strftime("%m-%d %H:%M") if etas else "완료")
        ws.append(row)
        r = ws.max_row
        for i, (p, _) in enumerate(cols, 4):
            v = ws.cell(r, i).value
            ws.cell(r, i).alignment = center
            if v == "—":
                ws.cell(r, i).fill = none_
            else:
                n, t = map(int, v.split("/"))
                if n and n == t:
                    ws.cell(r, i).fill = full
                elif n:
                    ws.cell(r, i).fill = part
        ws.cell(r, len(cols) + 4).alignment = center
        prev = s
    ws.freeze_panes = "D3"
    end = ws.max_row
    ws.append([])
    ws.append(["", "칠", "진하게 = 다 봤다 · 옅게 = 일부 · 흰 칸 = 아직 · 회색 = 그 강에 해당 산물 없음. 칸 = 검토를 거친 단위 / 그 강의 단위."])
    ws.append(["", "순서", f"검토는 산물 순서 {' → '.join(order)} 로 돈다(job.json review.order) — 한 산물을 다 보고 다음 산물로 넘어간다."])
    ws.append(["", "예상", f"남은 배치 {len(pend)}개 × 배치당 평균 {avg:.0f}초 기준. 사용 한도 대기·재검증(verify)·정오표 단계 시간은 넣지 않았다."])
    ws.append(["", "기준", "문제집은 연결된 요약·롱폼 항목이 가장 많은 장의 강에, 문제집 영상은 그 문항을 따라 붙였다(영상 낭독은 문항 검토에서 함께 본다)."
               + (f" 강에 붙이지 못한 단위: {', '.join(f'{PRODUCTS[k]} {v}' for k, v in unplaced.items())}." if unplaced else "")])
    for r in range(end + 2, ws.max_row + 1):
        ws.cell(r, 3).alignment = Alignment(wrap_text=False)


def openpyxl_col(i: int) -> str:
    from openpyxl.utils import get_column_letter
    return get_column_letter(i)


def run(job: Job) -> Path:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from errata.s03_review import _safe

    units = read_jsonl(job.p("00_대상", "units.jsonl"))
    machine = read_jsonl(job.p("02_기계검사", "findings.jsonl"))
    merged = {f["fid"]: f for f in read_jsonl(job.p("04_수용", "merged.jsonl"))}
    links: Dict[str, List[str]] = read_json(job.p("00_대상", "links.json"), {}) or {}
    src_of: Dict[str, str] = {}
    for q, ds in links.items():
        for d in ds:
            src_of.setdefault(d, q)

    # 검토 배치 결과
    rdir = job.p("03_검토")
    reviews: Dict[str, dict] = {}
    for p in sorted(rdir.glob("*.json")) if rdir.exists() else []:
        d = read_json(p, {}) or {}
        if d.get("batch"):
            reviews[d["batch"]] = d
    failed = {(read_json(p, {}) or {}).get("batch") for p in (rdir / "_실패").glob("*.json")} if (rdir / "_실패").exists() else set()

    # 단위별 상태
    by_uid_m: Dict[str, List[dict]] = collections.defaultdict(list)
    for f in machine:
        by_uid_m[f["uid"]].append(f)
    by_uid_r: Dict[str, List[dict]] = collections.defaultdict(list)
    clean, solve, seen_at, silent = set(), {}, {}, set()
    for b, d in reviews.items():
        for f in d.get("findings") or []:
            by_uid_r[f["uid"]].append(f)
        clean.update(d.get("clean") or [])
        silent.update(d.get("unmentioned") or [])
        for s in d.get("solves") or []:
            solve[s.get("uid")] = s
        for u in d.get("units") or []:
            seen_at[u] = d.get("at")

    def status(u: dict) -> str:
        b = u.get("batch")
        if not b:
            return ST_NONE
        if b not in reviews:
            return ST_FAIL if b in failed else ST_TODO
        if by_uid_r.get(u["uid"]):
            return ST_HIT
        if u["uid"] in clean:
            return ST_CLEAN
        return ST_SILENT

    def ts(x) -> str:
        return dt.datetime.fromtimestamp(x).strftime("%Y-%m-%d %H:%M") if x else ""

    def fsum(fs: List[dict]) -> str:
        return " / ".join(f"{f['severity']}·{f['type']}" + (f"({(merged.get(f['fid'], {}).get('verify') or {}).get('verdict')})"
                                                          if (merged.get(f['fid'], {}).get('verify') or {}).get('verdict') else "")
                          for f in fs)

    counts = collections.Counter((f["product"], f.get("rule") or f["origin"]) for f in machine)
    for fs in by_uid_r.values():
        for f in fs:
            counts[(f["product"], f.get("rule") or "review")] += 1
    units_by_p = collections.Counter(u["product"] for u in units)
    batched = {u["product"] for u in units if u.get("batch")}
    rules = machine_rules()

    wb = openpyxl.Workbook()
    wrap = Alignment(wrap_text=True, vertical="top")
    table = _table

    # 읽는 법
    ws = wb.active
    ws.title = "읽는 법"
    now = dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    st_cnt = collections.Counter(status(u) for u in units)
    lines = [
        [f"검수 기록 — {job.cfg.get('label', job.name)}", ""],
        ["갱신", now],
        ["쓰임", "산물에서 오류가 나왔을 때 「그 단위를 검수했나 · 언제 · 무엇을 기준으로 · 결과는」을 찾는다. "
                 "「단위별 기록」에서 위치·uid·제목으로 필터 → 프롬프트 기록 파일을 열면 검수단이 실제로 읽은 본문과 응답이 있다."],
        ["", ""],
        ["검토 상태", "뜻"],
        [ST_CLEAN, "검수단이 읽고 지적 없음으로 명시했다(clean)."],
        [ST_HIT, "검수단 지적이 1건 이상 있다. 「검토 지적」 열과 정오표 엑셀을 본다."],
        [ST_SILENT, "배치는 돌았는데 검수단이 이 단위를 지적에도 이상없음에도 적지 않았다 — 봤다고 볼 근거가 없다."],
        [ST_TODO, "이 단위의 배치가 아직 돌지 않았다."],
        [ST_FAIL, "배치 호출이 실패했다(03_검토/_실패). review 를 다시 돌리면 그것만 다시 한다."],
        [ST_NONE, "검수단 검토를 거치지 않는 산물(문제풀이 영상 V·덱 D) — 기계검사만. 「함께 본 원천」 열이 그 내용을 검토한 원천 단위다."],
        ["", ""],
        ["지금 현황", " · ".join(f"{k} {v}" for k, v in st_cnt.most_common())],
        ["한계", "검수단은 가상의 응답자다. 「이상없음」은 틀린 것이 없다는 보증이 아니라 「검토를 거쳤고 지적이 나오지 않았다」는 기록이다. "
                 "기계검사는 「검수 기준」 시트의 규칙에 걸리는 것만 잡는다."],
    ]
    for r in lines:
        ws.append(r)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 16, 120
    ws["A1"].font = Font(bold=True, size=13)
    ws["A5"].font = ws["B5"].font = Font(bold=True)
    for i in range(1, ws.max_row + 1):
        ws.cell(i, 2).alignment = wrap
        v = ws.cell(i, 1).value
        if v in ST_FILL:
            ws.cell(i, 1).fill = PatternFill("solid", fgColor=ST_FILL[v])

    # 검수 기준
    table(wb.create_sheet("검수 기준"),
          [("산물", 14), ("단계", 14), ("항목", 22), ("무엇을 보는가", 90), ("대상", 26), ("지금까지 지적", 12)],
          _criteria(job, units_by_p, batched, counts), wrap_cols=(4,))

    # 배치 기록
    allb: Dict[str, List[dict]] = collections.OrderedDict()
    for u in units:
        if u.get("batch"):
            allb.setdefault(u["batch"], []).append(u)
    brows = []
    for b, us in allb.items():
        d = reviews.get(b)
        uids = [u["uid"] for u in us]
        if d:
            fs = d.get("findings") or []
            sev = collections.Counter(f["severity"] for f in fs)
            st = "완료"
            cl = sum(1 for x in uids if x in clean and not by_uid_r.get(x))
            sl = sum(1 for u in us if status(u) == ST_SILENT)
        else:
            sev, fs, cl, sl = collections.Counter(), [], "", ""
            st = ST_FAIL if b in failed else ST_TODO
        pr = rdir / "_prompts" / f"{_safe(b)}.프롬프트.txt"
        brows.append([PRODUCTS[us[0]["product"]], b, f"{us[0]['loc']} ~ {us[-1]['loc']}" if len(us) > 1 else us[0]["loc"],
                      len(us), st, ts(d.get("at")) if d else "", d.get("sec", "") if d else "",
                      len(fs) if d else "", sev["상"] if d else "", sev["중"] if d else "", sev["하"] if d else "", cl, sl,
                      str(pr) if pr.exists() else "", str(pr.with_name(f"{_safe(b)}.응답.txt")) if pr.exists() else ""])
    table(wb.create_sheet("배치 기록"),
          [("산물", 14), ("배치", 18), ("범위", 60), ("단위", 7), ("상태", 8), ("검토 시각", 17), ("초", 6),
           ("지적", 7), ("상", 5), ("중", 5), ("하", 5), ("이상없음 단위", 10), ("언급 없음 단위", 10),
           ("프롬프트 기록", 50), ("응답 기록", 50)],
          brows, status_col=5, freeze="C2")

    # 단위별 기록
    rule_names = collections.defaultdict(list)
    for name, prods, _ in rules:
        for p in "QVSDLBT":
            if _rule_applies(prods, p):
                rule_names[p].append(name.replace("*", p))
    urows = []
    for u in units:
        p, uid = u["product"], u["uid"]
        st = status(u)
        mf, rf = by_uid_m.get(uid, []), by_uid_r.get(uid, [])
        s = solve.get(uid)
        solve_txt = ""
        if s:
            solve_txt = f"풀이 {s.get('solved')} · 표기 {s.get('marked')} · {'일치' if s.get('agree', True) else '불일치'}"
        src = src_of.get(uid, "")      # 파생 단위(V·S·L 예제 …) → 원천 문항
        urows.append([PRODUCTS[p], uid, u.get("loc", ""), (u.get("title") or "")[:80], u.get("batch") or "",
                      st, ts(seen_at.get(uid)), len(rf), fsum(rf), ", ".join(sorted({f.get('reviewer') or '' for f in rf} - {""})),
                      solve_txt, ", ".join(rule_names[p]), len(mf), fsum(mf), src, "", u.get("file", "")])
    U = {u["uid"]: u for u in units}
    for r in urows:
        if r[14] in U:
            r[15] = status(U[r[14]])
    table(wb.create_sheet("단위별 기록"),
          [("산물", 14), ("uid", 26), ("위치", 44), ("제목", 30), ("배치", 16), ("검토 상태", 12), ("검토 시각", 17),
           ("검토 지적", 7), ("검토 지적 내용(심각도·유형(재검증))", 40), ("지적한 검수자", 12), ("눈가림 풀이(문항)", 22),
           ("기계검사 항목", 40), ("기계 지적", 7), ("기계 지적 내용", 30), ("함께 본 원천", 14), ("원천 검토 상태", 12), ("파일", 60)],
          urows, status_col=6, freeze="C2")

    write_grid(wb, job, units, reviews, links)
    wb.active = 0
    out = job.p("06_검수표", "검수기록.xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(out)
    except PermissionError:
        out = out.with_name(f"검수기록_{dt.datetime.now():%m%d-%H%M}.xlsx")
        wb.save(out)
        warn(f"검수기록.xlsx 가 열려 있어 덮지 못했습니다 → {out.name} 로 저장")
    write_roster(job, collections.Counter((f["product"], f.get("reviewer") or "") for fs in by_uid_r.values() for f in fs),
                 batched, units_by_p)
    log(f"  검수 기록 → {out.name} · " + " · ".join(f"{k} {v}" for k, v in st_cnt.most_common()))
    return out
